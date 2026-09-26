import openpyxl
import datetime
import json
import os
import re

def parse_amount(val):
    if val is None or val == 'NA' or val == '':
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(',', '')
    if '+' in s:
        try:
            return sum(float(x.strip()) for x in s.split('+') if x.strip())
        except Exception:
            pass
    try:
        return float(s)
    except Exception:
        return 0.0

def categorize_reason(reason, default_type='expense'):
    r = (reason or '').lower().strip()
    if not r:
        return ('Other Expense', 'expense')
    
    if 'salary' in r:
        return ('Monthly Salary', 'income')
    if 'deposit' in r or 'income' in r:
        return ('Other Income', 'income')
    if 'sip' in r or 'motilal' in r or 'quant' in r or 'nippon' in r or 'mutual fund' in r:
        return ('Index Mutual Fund', 'expense')
    
    # Food & Dining
    if any(k in r for k in ['swiggy', 'zomato', 'blinkit', 'zepto', 'instamart', 'food', 'lunch', 'dinner', 'breakfast', 'tea', 'coffee', 'ice cream', 'cake', 'snacks', 'hotel', 'restaurant', 'biryani', 'pani puri', 'fried rice', 'panner', 'burger', 'pizza', 'chat', 'bakery', 'juice', 'sweet', 'egg', 'chicken', 'dosa', 'roti']):
        if 'swiggy' in r or 'zomato' in r:
            return ('Swiggy / Zomato', 'expense')
        if 'blinkit' in r or 'zepto' in r or 'instamart' in r:
            return ('Blinkit / Zepto / Instamart', 'expense')
        if 'coffee' in r or 'tea' in r or 'cafe' in r:
            return ('Coffee & Cafe', 'expense')
        if 'dinner' in r or 'lunch' in r or 'hotel' in r or 'restaurant' in r:
            return ('Dining Out & Restaurants', 'expense')
        return ('Groceries & Supermarket', 'expense')
    
    # Shopping
    if any(k in r for k in ['amazon', 'flipkart', 'myntra', 'shirt', 'pant', 'dress', 'shoe', 'shopping', 'cloth', 'electronics', 'gadget', 'bag', 'watch']):
        if 'amazon' in r:
            return ('Amazon Shopping', 'expense')
        if 'flipkart' in r:
            return ('Flipkart Shopping', 'expense')
        if 'myntra' in r or 'cloth' in r or 'shirt' in r or 'pant' in r:
            return ('Myntra & Clothing', 'expense')
        return ('Amazon Shopping', 'expense')
        
    # Travel & Commute
    if any(k in r for k in ['petrol', 'fuel', 'diesel', 'uber', 'ola', 'rapido', 'auto', 'metro', 'bus', 'train', 'irctc', 'flight', 'fastag', 'toll', 'bike', 'cab']):
        if 'petrol' in r or 'fuel' in r or 'diesel' in r:
            return ('Fuel (Petrol / Diesel)', 'expense')
        if 'uber' in r or 'ola' in r or 'rapido' in r or 'auto' in r or 'cab' in r:
            return ('Uber / Ola / Rapido', 'expense')
        if 'metro' in r or 'bus' in r:
            return ('Metro & Bus', 'expense')
        if 'train' in r or 'irctc' in r:
            return ('Train / IRCTC', 'expense')
        if 'fastag' in r or 'toll' in r:
            return ('Fastag & Toll', 'expense')
        return ('Travel & Commute', 'expense')
        
    # Entertainment & Sports
    if any(k in r for k in ['badminton', 'shuttle', 'court', 'movie', 'cinema', 'game', 'party', 'birthday', 'outing', 'trip', 'kencha', 'vidya mandir']):
        if 'badminton' in r or 'court' in r or 'shuttle' in r:
            return ('Gym & Fitness', 'expense')
        if 'movie' in r or 'cinema' in r:
            return ('Movies & Events', 'expense')
        if 'party' in r or 'birthday' in r or 'outing' in r:
            return ('Movies & Events', 'expense')
        return ('Movies & Events', 'expense')
        
    # Health & Medical
    if any(k in r for k in ['medicine', 'medical', 'doctor', 'hospital', 'tablet', 'pharmacy', 'clinic', 'dentist']):
        return ('Pharmacy & Medicines', 'expense')
        
    # Bills & Utilities
    if any(k in r for k in ['recharge', 'mobile', 'wifi', 'broadband', 'electricity', 'current', 'water', 'gas', 'cylinder', 'dth', 'currency', 'amma currency']):
        if 'recharge' in r or 'mobile' in r:
            return ('Mobile Recharge & Postpaid', 'expense')
        if 'wifi' in r or 'broadband' in r:
            return ('Broadband / WiFi', 'expense')
        if 'currency' in r:
            return ('Mobile Recharge & Postpaid', 'expense')
        return ('Electricity Bill', 'expense')
        
    # Housing & Pets
    if any(k in r for k in ['rent', 'room', 'house', 'maintenance', 'bekku', 'cat', 'dog', 'pet']):
        if 'bekku' in r or 'cat' in r or 'dog' in r or 'pet' in r:
            return ('Gifts & Celebrations', 'expense')
        return ('House Rent', 'expense')
        
    return ('Other Expense', 'expense')

def normalize_payment_method(pay_str):
    p = (pay_str or '').lower().strip()
    if not p:
        return 'Cash'
    if 'one card' in p or 'onecard' in p:
        return 'OneCard Credit Card'
    if 'gpay' in p or 'google pay' in p:
        return 'Google Pay (GPay)'
    if 'phone pay' in p or 'phonepe' in p or 'phone' in p:
        return 'PhonePe'
    if 'cred' in p:
        return 'CRED UPI'
    if 'paytm' in p:
        return 'Paytm UPI / Wallet'
    if 'amazon' in p:
        return 'Amazon Pay'
    if 'hdfc' in p:
        return 'HDFC Credit Card'
    if 'card' in p:
        return 'OneCard Credit Card'
    if 'cash' in p:
        return 'Cash'
    if 'online' in p or 'upi' in p:
        return 'Google Pay (GPay)'
    return 'Cash'

def extract_all():
    files = [
        (r'd:\OTHRES\mine only\2024_expenses.xlsx', 2024),
        (r'd:\OTHRES\mine only\2025_expenses.xlsx', 2025),
        (r'd:\OTHRES\mine only\2026_Expenses.xlsx', 2026),
    ]
    
    all_transactions = []
    
    for path, default_year in files:
        if not os.path.exists(path):
            print(f"Skipping missing file: {path}")
            continue
            
        wb = openpyxl.load_workbook(path, data_only=True)
        file_txns = []
        
        for s in wb.sheetnames:
            if s.lower() == 'sip':
                continue
            ws = wb[s]
            rows = list(ws.iter_rows(values_only=True))
            current_date = None
            current_month = None
            
            for r_idx, r in enumerate(rows):
                if not r or len(r) < 5:
                    continue
                c0, c1, c2, c3, c4 = r[0], r[1], r[2], r[3], r[4]
                
                if '2024' in path:
                    d_val = c1
                    amt_val = c3
                    pay_val = c4
                    reason_val = r[5] if len(r) > 5 else ''
                else:
                    d_val = c0
                    amt_val = c2
                    pay_val = c3
                    reason_val = c4
                    
                # Month header
                if isinstance(d_val, str) and d_val.strip().upper() in ['JANUARY', 'FEBRUARY', 'MARCH', 'APRIL', 'MAY', 'JUNE', 'JULY', 'AUGUST', 'SEPTEMBER', 'OCTOBER', 'NOVEMBER', 'DECEMBER']:
                    current_month = d_val.strip().title()
                    continue
                if isinstance(c0, str) and c0.strip().upper() in ['JANUARY', 'FEBRUARY', 'MARCH', 'APRIL', 'MAY', 'JUNE', 'JULY', 'AUGUST', 'SEPTEMBER', 'OCTOBER', 'NOVEMBER', 'DECEMBER']:
                    current_month = c0.strip().title()
                    
                if isinstance(d_val, (datetime.datetime, datetime.date)):
                    current_date = d_val.strftime('%Y-%m-%d')
                elif d_val == 'NA':
                    continue
                    
                amount = parse_amount(amt_val)
                if amount > 0 and current_date:
                    reason_str = str(reason_val or '').strip()
                    pay_str = str(pay_val or '').strip()
                    category_name, tx_type = categorize_reason(reason_str)
                    norm_payment = normalize_payment_method(pay_str)
                    
                    file_txns.append({
                        'date': current_date,
                        'amount': round(amount, 2),
                        'type': tx_type,
                        'raw_reason': reason_str,
                        'category': category_name,
                        'raw_payment': pay_str,
                        'payment_method': norm_payment,
                        'source_year': default_year
                    })
                    
        print(f"Extracted from {os.path.basename(path)}: {len(file_txns)} transactions, Total: Rs. {sum(t['amount'] for t in file_txns):,.2f}")
        all_transactions.extend(file_txns)
        
    print(f"\nTotal extracted across all years: {len(all_transactions)} transactions")
    print(f"Grand Total Amount: Rs. {sum(t['amount'] for t in all_transactions):,.2f}")
    
    # Save extracted JSON
    out_dir = r'C:\Users\ManojShankar\Desktop\Learnings\Expenses_website\Version_1.1\finance-tracker\src\data'
    os.makedirs(out_dir, exist_ok=True)
    out_file = os.path.join(out_dir, 'past_expenses.json')
    with open(out_file, 'w', encoding='utf-8') as f:
        json.dump(all_transactions, f, indent=2)
    print(f"Saved to {out_file}")

if __name__ == '__main__':
    extract_all()
